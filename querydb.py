import streamlit as st
import pandas as pd
from io import BytesIO
import datetime
import os
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
# import openpyxl
from openpyxl import load_workbook
import io
import msoffcrypto 
# from querydb import query_database
import pymysql 
import pandas as pd
from time import sleep

def query_database(account_number, query_date):
    try:
        # Establish connection to your database
        connection = pymysql.connect(host='192.168.15.197',
                                     user='ljbernas_bcp',
                                     password='$C4Ov9P52n1sh',
                                     database='bcrm',
                                     charset='utf8mb4',
                                     cursorclass=pymysql.cursors.DictCursor)

        # Prepare SQL query with placeholders for dynamic values
        sql = """
        SELECT
            client.client_name AS 'campaign',
            leads_result.leads_result_id AS 'ResultID',
            users.users_username AS 'Agent',
            leads.leads_chcode AS 'chCode',
            leads.leads_chname AS 'chName',
            leads.leads_placement AS 'placement',
            leads.leads_acctno AS 'AccountNumber',
            leads_status.leads_status_name AS 'Status',
            leads_result.leads_result_amount AS 'Amount',
            leads_result.leads_result_ts AS 'ResultDate',
            leads_result.leads_result_source AS 'source',
            leads.leads_endo_date AS 'EndoDate',
            leads.leads_ob AS 'OB',
            leads_result.leads_result_barcode_date
        FROM bcrm.leads_result
        LEFT JOIN bcrm.leads ON (leads_result.leads_result_lead = leads.leads_id)
        LEFT JOIN bcrm.client ON (leads.leads_client_id = client.client_id)
        LEFT JOIN bcrm.users ON (leads_result.leads_result_users = users.users_id)
        LEFT JOIN bcrm.users AS leads_users ON (leads_users.users_id = leads.leads_users_id)
        LEFT JOIN bcrm.leads_status ON (leads_result.leads_result_status_id = leads_status.leads_status_id)
        LEFT JOIN bcrm.leads_substatus ON (leads_result.leads_result_substatus_id = leads_substatus.leads_substatus_id) 
        WHERE
            leads_users.users_username <> 'POUT' 
            AND leads.leads_acctno = %s
            AND DATE(leads_result.leads_result_ts) = %s;
        """
        # Execute SQL query with parameters
        with connection.cursor() as cursor:
            cursor.execute(sql, (account_number, query_date))
            results = cursor.fetchall()

        return results[0] if results else None
        
    except Exception as e:
        print(f"Error querying database: {e}")
        sleep(5)
        query_database(account_number, query_date)

    finally:
        if connection:
            connection.close()